# -*- coding: utf-8 -*-
"""
Crea la base de datos SQLite (upgrade_ms.db) y la puebla a partir del
resumen maestro y los Excel de detalle. Idempotente: si la base ya existe
NO se sobreescribe salvo que se pase --reset.

Uso:
  python import_excel.py           # crea/puebla si esta vacia
  python import_excel.py --reset   # borra y reconstruye desde los Excel
"""
import openpyxl, glob, os, re, json, sys, datetime, sqlite3, unicodedata

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.dirname(os.path.abspath(__file__))
MASTER = os.path.join(BASE, "Resumen de comunicados_Actualizado_v1.xlsx")
DB = os.path.join(HERE, "upgrade_ms.db")

# --------- normalizacion de columnas ---------
def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[\s_\-\.]+", "", s)

SUB_NAME = ["subscriptionname", "suscripcion", "subscription"]
SUB_ID   = ["subscriptionid", "suscriptionid"]
RG       = ["grupoderecurso", "grupoderecursos", "resourcegroup", "resourcegroups"]
RES      = ["nombredelrecurso","resourcename","recursoafectado","recursoasociado",
            "functionappname","appservicename","storageaccountname","keyvaultname",
            "databricksname","servername","storageaccount","workspace","vmname",
            "account","recurso","name","nombre"]
GESTOR   = ["gestor"]
ESTADO   = ["estado","status","estado2"]

DONE_WORDS = ["completado","completo","migrado","migrada","cerrado","finalizado",
              "ok","resuelto","hecho","aplicado","revisado","done"]

def match_col(hn, cands, loose=True):
    for c in cands:
        for i, h in enumerate(hn):
            if h == c: return i
    if loose:
        for c in cands:
            for i, h in enumerate(hn):
                if c in h and len(h)-len(c) <= 4: return i
    return None

def to_str(v):
    if v is None: return ""
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

# --------- categorias amplias (consolidadas) ---------
CATEGORIAS = ["Compute", "Almacenamiento", "Redes", "Bases de datos",
              "Datos y Analítica", "Contenedores", "Seguridad e Identidad",
              "Gobernanza y Monitoreo", "FinOps y Reservas", "Otros"]

# El orden importa: la primera coincidencia gana.
_CAT_RULES = [
    (["databricks"], "Datos y Analítica"),
    (["reserv", "finops", "facturacion", "billing"], "FinOps y Reservas"),
    (["network", "cdn", "frontdoor", "vpn"], "Redes"),
    (["contenedor", "aks", "kubernetes", "container"], "Contenedores"),
    (["seguridad", "identidad", "entra", "rbac", "tls", "keyvault", "key vault", "sftp"], "Seguridad e Identidad"),
    (["gobernanza", "monitoreo", "monitor", "automation", "automatizacion",
      " ia", "/ia", "service health", "diagnostic"], "Gobernanza y Monitoreo"),
    (["base de datos", "mysql", "postgres", "redis", "cache", "sql"], "Bases de datos"),
    (["storage", "almacen", "disco", "disk", "blob"], "Almacenamiento"),
    (["compute", "app service", "runtime", "virtual desktop", "desktop",
      "gpu", "api", "python", "node", "windows app", "vm"], "Compute"),
]

def map_cat(raw):
    """Convierte una categoria original en una de las categorias amplias."""
    if not raw: return ""
    s = unicodedata.normalize("NFKD", str(raw))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    for keys, cat in _CAT_RULES:
        if any(k in s for k in keys): return cat
    return "Otros"

def header_row(rows):
    for i, r in enumerate(rows[:8]):
        if len([c for c in r if c not in (None, "")]) >= 2: return i
    return 0

def build_file_index():
    idx = {}
    for f in glob.glob(os.path.join(BASE, "*.xlsx")):
        b = os.path.basename(f)
        if b.startswith("Resumen de comunicados"): continue
        idx[norm(re.sub(r"^\d+[_\-\s]*", "", b))] = f
        idx[norm(b)] = f
    return idx

def parse_date(s):
    """devuelve ISO YYYY-MM-DD o None"""
    if not s: return None
    return s if re.match(r"^\d{4}-\d{2}-\d{2}$", s) else None

# --------- base de clientes (cliente -> suscripciones) ---------
CSP = os.path.join(BASE, "CSP - GESTION DE CLIENTES.xlsx")

def build_cliente_rows():
    """Devuelve [(cliente, tenant_id, suscripcion, sub_id), ...] de la hoja CLIENTES.
    CLIENTE y TENANT ID solo figuran en la 1a fila de cada cliente: se arrastran hacia abajo."""
    out = []
    if not os.path.exists(CSP):
        return out
    wb = openpyxl.load_workbook(CSP, read_only=True, data_only=True)
    ws = wb["CLIENTES"] if "CLIENTES" in wb.sheetnames else None
    if ws is None:
        wb.close(); return out
    rows = list(ws.iter_rows(values_only=True))
    hn = [norm(x) for x in rows[0]]
    ic = hn.index("cliente") if "cliente" in hn else 1
    iten = next((i for i, h in enumerate(hn) if h == "tenantid"), 3)
    isu = next((i for i, h in enumerate(hn) if h == "suscripcion"), 8)
    iid = next((i for i, h in enumerate(hn) if "iddelasuscripcion" in h or h == "idsuscripcion"), 9)
    cur = cur_ten = ""
    for r in rows[1:]:
        cli = to_str(r[ic]) if ic < len(r) else ""
        if cli:
            cur = cli
            cur_ten = to_str(r[iten]) if iten < len(r) else ""
        su = to_str(r[isu]) if isu < len(r) else ""
        idd = to_str(r[iid]) if iid < len(r) else ""
        if cur and (su or idd):
            out.append((cur, cur_ten, su, idd))
    wb.close()
    return out

def build_cliente_map():
    """(por_nombre, por_id) suscripcion->cliente, a partir de las filas de la hoja."""
    by_name, by_id = {}, {}
    for cli, ten, su, idd in build_cliente_rows():
        if su: by_name[norm(su)] = cli
        if idd: by_id[idd.lower()] = cli
    return by_name, by_id

def cliente_de(name, sid, by_name, by_id):
    c = by_name.get(norm(name)) or by_id.get((sid or "").strip().lower())
    if c: return c
    nl = (name or "").lower()
    if "g&s" in nl or "g and s" in nl or nl.startswith("azure g"): return "G&S (interno)"
    return ""

# --------- esquema ---------
SCHEMA = """
CREATE TABLE comunicados(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  n TEXT, titulo TEXT NOT NULL, categoria TEXT,
  fecha_recepcion TEXT, fecha_limite TEXT,
  resumen TEXT, fuente TEXT, archivo TEXT,
  estado TEXT, responsable TEXT,
  ultima_actualizacion TEXT, proxima_actualizacion TEXT,
  observaciones TEXT,
  kql TEXT,                      -- query KQL de Resource Graph usado para inventariar
  origen TEXT DEFAULT 'excel',
  created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE recursos(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  comunicado_id INTEGER NOT NULL REFERENCES comunicados(id) ON DELETE CASCADE,
  hoja TEXT,
  cliente TEXT,
  suscripcion TEXT, suscripcion_id TEXT,
  grupo_recurso TEXT, nombre_recurso TEXT,
  gestor TEXT, estado TEXT,
  extra TEXT,                    -- JSON con columnas restantes
  revisado INTEGER DEFAULT 0,
  revisado_por TEXT, revisado_at TEXT,
  notas TEXT,
  created_at TEXT DEFAULT (datetime('now')),   -- fecha de inventariado del recurso
  inventario_id INTEGER
);
CREATE TABLE inventarios(          -- lotes de inventario: cada corrida del query
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  comunicado_id INTEGER NOT NULL REFERENCES comunicados(id) ON DELETE CASCADE,
  fecha TEXT, kql TEXT, nota TEXT,
  created_at TEXT DEFAULT (datetime('now'))
);
CREATE INDEX ix_inv_com ON inventarios(comunicado_id);
CREATE INDEX ix_rec_com ON recursos(comunicado_id);
CREATE INDEX ix_rec_sub ON recursos(suscripcion);
CREATE INDEX ix_rec_rev ON recursos(revisado);

CREATE TABLE clientes(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  nombre TEXT UNIQUE NOT NULL,
  ext_id TEXT,
  created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE suscripciones(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  cliente_id INTEGER NOT NULL REFERENCES clientes(id) ON DELETE CASCADE,
  nombre TEXT, sub_id TEXT,
  created_at TEXT DEFAULT (datetime('now'))
);
CREATE INDEX ix_sus_cli ON suscripciones(cliente_id);
"""

CLIENTES_DDL = """
CREATE TABLE IF NOT EXISTS clientes(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  nombre TEXT UNIQUE NOT NULL,
  ext_id TEXT,
  created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS suscripciones(
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  cliente_id INTEGER NOT NULL REFERENCES clientes(id) ON DELETE CASCADE,
  nombre TEXT, sub_id TEXT,
  created_at TEXT DEFAULT (datetime('now'))
);
CREATE INDEX IF NOT EXISTS ix_sus_cli ON suscripciones(cliente_id);
"""

def seed_clientes(conn):
    """Puebla clientes/suscripciones desde la hoja CLIENTES (con fill-down)."""
    cur = conn.cursor()
    cli_id = {}
    n_c = n_s = 0
    for cliente, tenant, su, sid in build_cliente_rows():
        if cliente not in cli_id:
            cur.execute("INSERT OR IGNORE INTO clientes(nombre,ext_id) VALUES(?,?)", (cliente, tenant))
            cli_id[cliente] = cur.execute("SELECT id FROM clientes WHERE nombre=?", (cliente,)).fetchone()[0]
            n_c += 1
        if su or sid:
            cur.execute("INSERT INTO suscripciones(cliente_id,nombre,sub_id) VALUES(?,?,?)",
                        (cli_id[cliente], su, sid))
            n_s += 1
    conn.commit()
    return n_c, n_s

def process_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return None
    hidx = header_row(rows)
    header = list(rows[hidx])
    last = len(header)
    while last > 0 and to_str(header[last-1]) == "": last -= 1
    header = header[:last]
    hstr = [to_str(h) for h in header]
    hn = [norm(h) for h in header]
    i_subn, i_subid = match_col(hn, SUB_NAME), match_col(hn, SUB_ID)
    i_rg, i_res = match_col(hn, RG), match_col(hn, RES)
    i_gestor, i_estado = match_col(hn, GESTOR, loose=False), match_col(hn, ESTADO, loose=False)
    used = {i for i in (i_subn,i_subid,i_rg,i_res,i_gestor,i_estado) if i is not None}
    out = []
    for r in rows[hidx+1:]:
        cells = [to_str(c) for c in list(r)[:last]]
        if not any(cells): continue
        extra = {hstr[i]: cells[i] for i in range(last) if i not in used and cells[i]}
        estado = cells[i_estado] if i_estado is not None and i_estado < len(cells) else ""
        revisado = 1 if any(w in estado.lower() for w in DONE_WORDS) else 0
        out.append({
            "hoja": ws.title,
            "suscripcion": cells[i_subn] if i_subn is not None else "",
            "suscripcion_id": cells[i_subid] if i_subid is not None else "",
            "grupo_recurso": cells[i_rg] if i_rg is not None else "",
            "nombre_recurso": cells[i_res] if i_res is not None else "",
            "gestor": cells[i_gestor] if i_gestor is not None else "",
            "estado": estado,
            "extra": json.dumps(extra, ensure_ascii=False),
            "revisado": revisado,
        })
    score = (1 if i_res is not None else 0) + (1 if i_rg is not None else 0) + \
            (1 if (i_subn is not None or i_subid is not None) else 0)
    return {"rows": out, "score": score*100000 + len(out)}

def best_sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        s = process_sheet(ws)
        if s and s["rows"]: sheets.append(s)
    wb.close()
    sheets.sort(key=lambda x: x["score"], reverse=True)
    return sheets

def seed(conn):
    fidx = build_file_index()
    by_name, by_id = build_cliente_map()
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb["Resumen de Comunicados"]
    rows = list(ws.iter_rows(values_only=True))
    cur = conn.cursor()
    seed_ts = datetime.datetime.now().isoformat(timespec="seconds")
    n_com = n_rec = 0
    for r in rows[1:]:
        if r[1] is None: continue
        archivo = to_str(r[7])
        cur.execute("""INSERT INTO comunicados
          (n,titulo,categoria,fecha_recepcion,fecha_limite,resumen,fuente,archivo,
           estado,responsable,ultima_actualizacion,proxima_actualizacion,observaciones,origen)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?, 'excel')""",
          (to_str(r[0]), to_str(r[1]), map_cat(to_str(r[2])), parse_date(to_str(r[3])),
           parse_date(to_str(r[4])), to_str(r[5]), to_str(r[6]), archivo, to_str(r[8]),
           to_str(r[9]), parse_date(to_str(r[10])), parse_date(to_str(r[11])), to_str(r[12])))
        cid = cur.lastrowid; n_com += 1
        mf = None
        if archivo:
            k = norm(re.sub(r"^\d+[_\-\s]*", "", archivo))
            mf = fidx.get(k) or fidx.get(norm(archivo))
        if mf:
            inv = cur.execute(
                "INSERT INTO inventarios(comunicado_id,fecha,kql,nota,created_at) VALUES(?,?,?,?,?)",
                (cid, seed_ts, "", "Carga inicial", seed_ts)).lastrowid
            got = 0
            for sh in best_sheets(mf):
                for row in sh["rows"]:
                    cli = cliente_de(row["suscripcion"], row["suscripcion_id"], by_name, by_id)
                    cur.execute("""INSERT INTO recursos
                      (comunicado_id,hoja,cliente,suscripcion,suscripcion_id,grupo_recurso,
                       nombre_recurso,gestor,estado,extra,revisado,created_at,inventario_id)
                      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                      (cid, row["hoja"], cli, row["suscripcion"], row["suscripcion_id"],
                       row["grupo_recurso"], row["nombre_recurso"], row["gestor"],
                       row["estado"], row["extra"], row["revisado"], seed_ts, inv))
                    n_rec += 1; got += 1
            if not got:
                cur.execute("DELETE FROM inventarios WHERE id=?", (inv,))
    wb.close()
    conn.commit()
    nc, ns = seed_clientes(conn)
    print(f"Sembrado: {n_com} comunicados, {n_rec} recursos, {nc} clientes, {ns} suscripciones.")

def migrate_clientes():
    """Sobre la base existente: agrega columna cliente (si falta) y la puebla
    desde la base de clientes, sin borrar datos."""
    conn = sqlite3.connect(DB)
    cols = [r[1] for r in conn.execute("PRAGMA table_info(recursos)").fetchall()]
    if "cliente" not in cols:
        conn.execute("ALTER TABLE recursos ADD COLUMN cliente TEXT")
        print("Columna 'cliente' agregada.")
    by_name, by_id = build_cliente_map()
    if not by_name:
        print("No se encontro la hoja CLIENTES en", CSP); conn.close(); return
    subs = conn.execute("SELECT DISTINCT suscripcion, suscripcion_id FROM recursos").fetchall()
    n = 0
    for name, sid in subs:
        cli = cliente_de(name or "", sid or "", by_name, by_id)
        if cli:
            conn.execute("UPDATE recursos SET cliente=? WHERE suscripcion=?", (cli, name))
            n += 1
    conn.commit()
    # tablas editables clientes/suscripciones
    conn.executescript(CLIENTES_DDL)
    ccols = [r[1] for r in conn.execute("PRAGMA table_info(clientes)").fetchall()]
    if "ext_id" not in ccols:
        conn.execute("ALTER TABLE clientes ADD COLUMN ext_id TEXT")
        print("Columna 'ext_id' (id de cliente) agregada.")
    if conn.execute("SELECT COUNT(*) FROM clientes").fetchone()[0] == 0:
        nc, ns = seed_clientes(conn)
        print(f"Tablas clientes/suscripciones sembradas: {nc} clientes, {ns} suscripciones.")
    else:  # poblar ext_id (tenant id) de clientes existentes desde el Excel
        ten = {}
        for cliente, tenant, su, sid in build_cliente_rows():
            if tenant and cliente not in ten: ten[cliente] = tenant
        for cliente, tenant in ten.items():
            conn.execute("UPDATE clientes SET ext_id=? WHERE nombre=? AND (ext_id IS NULL OR ext_id='')",
                         (tenant, cliente))
        conn.commit()
    tot = conn.execute("SELECT COUNT(*) FROM recursos WHERE cliente<>''").fetchone()[0]
    ncli = conn.execute("SELECT COUNT(DISTINCT cliente) FROM recursos WHERE cliente<>''").fetchone()[0]
    conn.close()
    print(f"Migracion OK: {n} suscripciones mapeadas, {tot} recursos con cliente, {ncli} clientes.")

def main():
    if "--clientes" in sys.argv:
        migrate_clientes(); return
    reset = "--reset" in sys.argv
    if os.path.exists(DB):
        if reset:
            os.remove(DB); print("Base anterior eliminada.")
        else:
            print("La base ya existe. Usa --reset para reconstruir, o --clientes para mapear clientes."); return
    conn = sqlite3.connect(DB)
    conn.executescript(SCHEMA)
    seed(conn)
    conn.close()
    print(f"Base creada en {DB}")

if __name__ == "__main__":
    main()
