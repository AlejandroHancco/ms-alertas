# -*- coding: utf-8 -*-
"""Inventarios (lotes), recursos afectados, importación CSV/XLSX y revisión."""
import csv
import io
import json

from . import db, parsing
from .clientes import ClienteMap


# ---------------- inventarios (lotes) ----------------
def list_inventarios(cid):
    # mismo orden que define el "último lote" (created_at): el primero es el vigente
    return db.rows("""
      SELECT i.*,
        (SELECT COUNT(*) FROM recursos r WHERE r.inventario_id=i.id) AS n_recursos,
        (SELECT COUNT(*) FROM recursos r WHERE r.inventario_id=i.id AND r.revisado=1) AS n_revisados,
        (SELECT COUNT(DISTINCT r.cliente) FROM recursos r WHERE r.inventario_id=i.id AND COALESCE(r.cliente,'')<>'') AS n_clientes,
        (SELECT COUNT(DISTINCT r.suscripcion) FROM recursos r WHERE r.inventario_id=i.id AND COALESCE(r.suscripcion,'')<>'') AS n_subs
      FROM inventarios i WHERE i.comunicado_id=%s ORDER BY i.created_at DESC, i.id DESC""", (cid,))


def add_inventario(cid, data):
    """Crea un lote vacío (0 recursos): registra que, a esa fecha, no hay recursos afectados."""
    with db.tx() as con:
        if not con.execute("SELECT 1 FROM comunicados WHERE id=%s", (cid,)).fetchone():
            raise ValueError("Comunicado no existe.")
        ts = db.now()
        inv = con.execute("INSERT INTO inventarios(comunicado_id,fecha,kql,created_at) VALUES(%s,%s,%s,%s) RETURNING id",
                          (cid, ts, data.get("kql") or "", ts)).fetchone()["id"]
    return {"ok": True, "inventario_id": inv, "importados": 0}


def update_inventario(iid, data):
    cols = [f for f in ("kql", "fecha") if f in data]
    if cols:
        db.run(f"UPDATE inventarios SET {','.join(f + '=%s' for f in cols)} WHERE id=%s",
               [data[f] for f in cols] + [iid])
    return {"ok": True}


def delete_inventario(iid):
    """Elimina un lote y los recursos que trajo."""
    with db.tx() as con:
        n = con.execute("DELETE FROM recursos WHERE inventario_id=%s", (iid,)).rowcount
        con.execute("DELETE FROM inventarios WHERE id=%s", (iid,))
    return {"ok": True, "recursos_eliminados": n}


# ---------------- recursos ----------------
def get_recursos(cid):
    out = db.rows("SELECT * FROM recursos WHERE comunicado_id=%s ORDER BY suscripcion,grupo_recurso,nombre_recurso", (cid,))
    for d in out:
        try:
            d["extra"] = json.loads(d["extra"]) if d["extra"] else {}
        except ValueError:
            d["extra"] = {}
    return out


def _read_table(ext, raw):
    """Devuelve (header, rows) desde CSV o XLSX, recortando columnas vacías al final."""
    ext = (ext or "").lower().lstrip(".")
    if ext in ("csv", "txt"):
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:3000]
        delim = ";" if sample.count(";") > sample.count(",") else ","
        rows = [[parsing.to_str(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    elif ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        best = max(wb.worksheets, key=lambda ws: ws.max_row or 0)   # la hoja con más filas
        rows = [[parsing.to_str(c) for c in r] for r in best.iter_rows(values_only=True)]
        wb.close()
    else:
        raise ValueError("Formato no soportado: usa .csv o .xlsx")
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("El archivo está vacío.")
    hidx = parsing.header_row(rows)
    header = rows[hidx]
    last = len(header)
    while last > 0 and header[last - 1].strip() == "":
        last -= 1
    return header[:last], [r[:last] for r in rows[hidx + 1:]]


def import_recursos(cid, ext, raw, replace=False, kql=""):
    header, data = _read_table(ext, raw)
    hn = [parsing.norm(h) for h in header]
    i_subn = parsing.match_col(hn, parsing.SUB_NAME)
    i_subid = parsing.match_col(hn, parsing.SUB_ID)
    i_rg = parsing.match_col(hn, parsing.RG)
    i_res = parsing.match_col(hn, parsing.RES)
    i_gestor = parsing.match_col(hn, parsing.GESTOR, loose=False)
    i_estado = parsing.match_col(hn, parsing.ESTADO, loose=False)

    missing = []
    if i_subn is None and i_subid is None:
        missing.append("Suscripción")
    if i_rg is None:
        missing.append("Grupo de Recurso (RG)")
    if i_res is None:
        missing.append("Nombre del Recurso")
    if missing:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(missing) +
                         ". Columnas detectadas en el archivo: " + ", ".join(h for h in header if h))

    with db.tx() as con:
        if not con.execute("SELECT 1 FROM comunicados WHERE id=%s", (cid,)).fetchone():
            raise ValueError("Comunicado no existe.")
        cliente_para = ClienteMap(con)
        if replace:   # empezar de cero: borra recursos y lotes anteriores
            con.execute("DELETE FROM recursos WHERE comunicado_id=%s", (cid,))
            con.execute("DELETE FROM inventarios WHERE comunicado_id=%s", (cid,))
        ts = db.now()   # un solo sello para todo el lote = un evento de inventariado
        inv = con.execute("INSERT INTO inventarios(comunicado_id,fecha,kql,created_at) VALUES(%s,%s,%s,%s) RETURNING id",
                          (cid, ts, kql or "", ts)).fetchone()["id"]
        filas = []
        for r in data:
            def cell(i):
                return r[i] if (i is not None and i < len(r)) else ""
            res = cell(i_res)
            if not (res or cell(i_rg) or cell(i_subn)):
                continue
            # fila original completa, tal cual el archivo subido (todas las columnas, en orden)
            extra = {header[i]: (r[i] if i < len(r) else "") for i in range(len(header)) if header[i].strip()}
            estado = cell(i_estado)
            revisado = 1 if any(w in estado.lower() for w in parsing.DONE_WORDS) else 0
            filas.append((cid, cliente_para(cell(i_subn), cell(i_subid)), cell(i_subn), cell(i_subid),
                          cell(i_rg), res, cell(i_gestor), estado,
                          json.dumps(extra, ensure_ascii=False), revisado, ts, inv))
        if filas:
            with con.cursor() as cur:
                cur.executemany("""INSERT INTO recursos
                  (comunicado_id,hoja,cliente,suscripcion,suscripcion_id,grupo_recurso,nombre_recurso,
                   gestor,estado,extra,revisado,created_at,inventario_id)
                  VALUES(%s,'importado',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", filas)
        else:   # nada válido: no dejes un lote vacío
            con.execute("DELETE FROM inventarios WHERE id=%s", (inv,))
    return {"ok": True, "importados": len(filas), "reemplazado": replace, "inventario_id": inv,
            "columnas": {"suscripcion": header[i_subn] if i_subn is not None else header[i_subid],
                         "rg": header[i_rg], "recurso": header[i_res]}}


# ---------------- revisión ----------------
def patch_recurso(rid, data, who):
    sets, vals = [], []
    for f in ("estado", "gestor", "notas"):
        if f in data:
            sets.append(f + "=%s"); vals.append(data[f])
    if "revisado" in data:
        val = 1 if data["revisado"] else 0
        sets += ["revisado=%s", "revisado_at=%s", "revisado_por=%s"]
        vals += [val, db.now() if val else None, who if val else None]
    if sets:
        db.run(f"UPDATE recursos SET {','.join(sets)} WHERE id=%s", vals + [rid])
    return {"ok": True}


def bulk_review(data, who):
    """Marca revisado por lista de ids, o por comunicado (+ suscripción opcional)."""
    val = 1 if data.get("revisado", True) else 0
    at, por = (db.now(), who) if val else (None, None)
    sets = "revisado=%s, revisado_at=%s, revisado_por=%s"
    n = 0
    if data.get("ids"):
        n = db.run(f"UPDATE recursos SET {sets} WHERE id = ANY(%s)", (val, at, por, [int(i) for i in data["ids"]]))
    elif "comunicado_id" in data:
        if data.get("suscripcion"):
            n = db.run(f"UPDATE recursos SET {sets} WHERE comunicado_id=%s AND suscripcion=%s",
                       (val, at, por, data["comunicado_id"], data["suscripcion"]))
        else:
            n = db.run(f"UPDATE recursos SET {sets} WHERE comunicado_id=%s", (val, at, por, data["comunicado_id"]))
    return {"ok": True, "actualizados": n}
