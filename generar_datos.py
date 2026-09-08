# -*- coding: utf-8 -*-
"""
Generador de datos para la Plataforma de Comunicados (Upgrade MS).
Lee el resumen maestro y cada Excel de detalle, normaliza las columnas
clave (Suscripcion / Grupo de Recurso / Nombre del Recurso) y produce
un archivo data.js autocontenido que consume index.html.

Uso:  python generar_datos.py
"""
import openpyxl, glob, os, json, re, unicodedata, datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = os.path.join(BASE, "Resumen de comunicados_Actualizado_v1.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.js")


def norm(s):
    """minusculas, sin acentos, sin espacios ni separadores."""
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r"[\s_\-\.]+", "", s)
    return s


# Listas de sinonimos (ya normalizados) en orden de prioridad
SUB_NAME = ["subscriptionname", "suscripcion", "subscription"]
SUB_ID = ["subscriptionid", "suscriptionid", "subscriptionidid"]
RG = ["grupoderecurso", "grupoderecursos", "resourcegroup", "resourcegroups"]
RES = [
    "nombredelrecurso", "resourcename", "recursoafectado", "recursoasociado",
    "functionappname", "appservicename", "storageaccountname", "keyvaultname",
    "databricksname", "servername", "storageaccount", "workspace",
    "vmname", "account", "recurso", "name", "nombre",
]
# Columna "ID" suelta (en algunos archivos es el id de suscripcion)
SUB_ID_LOOSE = ["id"]


def match_col(headers_norm, candidates):
    """Devuelve el indice de la primera cabecera que coincida con algun candidato."""
    for cand in candidates:
        for i, h in enumerate(headers_norm):
            if h == cand:
                return i
    # coincidencia parcial (contiene)
    for cand in candidates:
        for i, h in enumerate(headers_norm):
            if cand in h and len(h) - len(cand) <= 4:
                return i
    return None


def to_str(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def find_header_row(rows):
    """Primera fila con >=2 celdas no vacias."""
    for i, r in enumerate(rows[:8]):
        nn = [c for c in r if c not in (None, "")]
        if len(nn) >= 2:
            return i
    return 0


def process_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None
    hidx = find_header_row(rows)
    header = rows[hidx]
    # recortar columnas vacias al final
    last = len(header)
    while last > 0 and (header[last - 1] is None or str(header[last - 1]).strip() == ""):
        last -= 1
    header = list(header[:last])
    headers_str = [to_str(h) for h in header]
    headers_norm = [norm(h) for h in header]

    i_subn = match_col(headers_norm, SUB_NAME)
    i_subid = match_col(headers_norm, SUB_ID)
    if i_subid is None:
        i_subid = match_col(headers_norm, SUB_ID_LOOSE)
    i_rg = match_col(headers_norm, RG)
    i_res = match_col(headers_norm, RES)

    data_rows = []
    for r in rows[hidx + 1:]:
        cells = [to_str(c) for c in list(r)[:last]]
        if not any(cells):
            continue
        data_rows.append(cells)

    key = {
        "subName": i_subn,
        "subId": i_subid,
        "rg": i_rg,
        "res": i_res,
    }
    # score: cuantas de las 3 claves obligatorias tiene + nro de filas
    have = sum(1 for k in ("res", "rg") if key[k] is not None)
    have += 1 if (i_subn is not None or i_subid is not None) else 0
    return {
        "name": ws.title,
        "headers": headers_str,
        "keys": key,
        "rows": data_rows,
        "score": have * 10000 + len(data_rows),
        "count": len(data_rows),
    }


def build_file_index():
    """mapa: nombre normalizado (sin prefijo NN_) -> ruta real."""
    idx = {}
    for f in glob.glob(os.path.join(BASE, "*.xlsx")):
        base = os.path.basename(f)
        if base.startswith("Resumen de comunicados"):
            continue
        stripped = re.sub(r"^\d+[_\-\s]*", "", base)
        idx[norm(stripped)] = f
        idx[norm(base)] = f
    return idx


def process_detail_file(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        s = process_sheet(ws)
        if s and s["count"] > 0:
            sheets.append(s)
    wb.close()
    sheets.sort(key=lambda x: x["score"], reverse=True)
    return sheets


def main():
    fidx = build_file_index()
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb["Resumen de Comunicados"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]

    comunicados = []
    for r in rows[1:]:
        if r[1] is None:
            continue
        archivo = to_str(r[7])
        detalle = []
        matched_file = None
        if archivo:
            key = norm(re.sub(r"^\d+[_\-\s]*", "", archivo))
            matched_file = fidx.get(key) or fidx.get(norm(archivo))
        if matched_file:
            try:
                detalle = process_detail_file(matched_file)
            except Exception as e:
                print("ERROR leyendo", matched_file, e)
        comunicados.append({
            "n": to_str(r[0]),
            "titulo": to_str(r[1]),
            "categoria": to_str(r[2]),
            "fechaRecepcion": to_str(r[3]),
            "fechaLimite": to_str(r[4]),
            "resumen": to_str(r[5]),
            "fuente": to_str(r[6]),
            "archivo": archivo,
            "estado": to_str(r[8]),
            "responsable": to_str(r[9]),
            "ultimaActualizacion": to_str(r[10]),
            "proximaActualizacion": to_str(r[11]),
            "observaciones": to_str(r[12]),
            "archivoEncontrado": bool(matched_file),
            "sheets": detalle,
        })
    wb.close()

    payload = {
        "generado": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "totalComunicados": len(comunicados),
        "comunicados": comunicados,
    }
    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write("window.DATA = ")
        json.dump(payload, fh, ensure_ascii=False, indent=1)
        fh.write(";")
    con_recursos = sum(1 for c in comunicados if c["sheets"])
    print(f"OK -> {OUT}")
    print(f"Comunicados: {len(comunicados)} | con recursos: {con_recursos}")


if __name__ == "__main__":
    main()
